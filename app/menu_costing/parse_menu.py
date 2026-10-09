"""Read a menu Sanjoy sent. Item and price columns are required. A blank price stays blank."""

import csv
import io
from pathlib import Path

from app.procurement.numbers import parse_number

ITEM_HEADERS = {
    "item",
    "item name",
    "menu item",
    "dish",
    "name",
    "product",
    "menu",
}
PRICE_HEADERS = {"price", "rate", "mrp", "selling price", "menu price"}
CATEGORY_HEADERS = {"category", "section", "group", "course"}


def parse_upload(filename, payload):
    suffix = Path(filename or "").suffix.casefold()
    if suffix == ".csv":
        return _from_rows(_csv_rows(payload))
    if suffix in {".xlsx", ".xlsm"}:
        return _from_rows(_xlsx_rows(payload))
    if suffix == ".xls":
        return {"items": [], "error": "Save this menu as .xlsx or .csv and upload it again."}
    if suffix == ".pdf":
        return _pdf_result(payload)
    return {"items": [], "error": "Upload an Excel, CSV, or PDF menu."}


def parse_menu_text(text):
    """Rows from extracted PDF text. Used when a PDF has a table as text."""
    lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
    rows = []
    for line in lines:
        if "," in line:
            rows.append([cell.strip() for cell in line.split(",")])
        else:
            parts = [part.strip() for part in line.split("  ") if part.strip()]
            if len(parts) < 2:
                parts = line.split()
            rows.append(parts)
    return _from_rows(rows)


def _csv_rows(payload):
    text = payload.decode("utf-8-sig", errors="replace")
    sample = text[:2000]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _xlsx_rows(payload):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        sheet = book.active
        return [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()


def _pdf_result(payload):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(payload))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return parse_menu_text(text)


def _from_rows(rows):
    header_at = None
    item_at = price_at = category_at = None
    for index, row in enumerate(rows[:25]):
        cells = [_header(cell) for cell in row]
        item_at = _find(cells, ITEM_HEADERS)
        price_at = _find(cells, PRICE_HEADERS)
        if item_at is None or price_at is None:
            continue
        category_at = _find(cells, CATEGORY_HEADERS)
        header_at = index
        break
    if header_at is None:
        return {
            "items": [],
            "error": "No item and price columns were found. The menu was not saved.",
        }
    found = {}
    for row in rows[header_at + 1 :]:
        if item_at >= len(row):
            continue
        name = str(row[item_at] or "").strip()
        if not name or _header(name) in ITEM_HEADERS:
            continue
        price_cell = row[price_at] if price_at < len(row) else None
        category = ""
        if category_at is not None and category_at < len(row) and row[category_at] is not None:
            category = str(row[category_at]).strip()
        found[name.casefold()] = {
            "item": name,
            "category": category,
            "price": parse_number(price_cell),
        }
    items = list(found.values())
    if not items:
        return {"items": [], "error": "The menu had a header and no items, so it was not saved."}
    return {"items": items, "error": ""}


def _header(value):
    return " ".join(str(value or "").strip().casefold().split())


def _find(cells, names):
    for index, cell in enumerate(cells):
        if cell in names:
            return index
    return None
