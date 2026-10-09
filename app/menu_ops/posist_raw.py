"""Read Tony's per-store Posist folders into menu-mix and channel rows.

Drop a new store in as one folder. The folder name is the store code
(0004, 002, 01-0001, 02-0013). A folder may be partial: menu workbooks,
a channel-range file, and a daily file are each optional. store.txt is
the Posist label when the folder name alone would match two stores.
"""

import csv
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from app.store_health.contract import parse_date, parse_number
from app.store_health.stores import _name_and_code, match_menu_store

MENU_OUT = "menu_mix_{start}_{end}.csv"
CHANNEL_OUT = "channel_sales_{start}_{end}.csv"
_MENU_FILE = re.compile(
    r"menu_items_(\d{4}-\d{2}-\d{2})_to_(\d{4}-\d{2}-\d{2}|\d{2}-\d{2})",
    re.IGNORECASE,
)


def posist_raw_directory(root=None):
    if root:
        return Path(root)
    return Path(__file__).resolve().parents[2] / "data" / "posist_raw"


def load_posist_raw(raw_root, posist_labels):
    """Return menu rows, channel rows, and a per-store coverage list."""
    raw_root = Path(raw_root)
    labels = [str(label).strip() for label in posist_labels if str(label).strip()]
    menu_rows = []
    channel_rows = []
    covered = []
    warnings = []
    if not raw_root.exists():
        warnings.append(f"{raw_root} is not on file.")
        return {"menu_rows": [], "channel_rows": [], "stores": [], "warnings": warnings}
    for folder in sorted(path for path in raw_root.iterdir() if path.is_dir()):
        store, why = _store_name(folder, labels)
        if not store:
            warnings.append(f"{folder.name} was skipped: {why}")
            continue
        menu, menu_warnings = _menu_rows(folder, store)
        channels, channel_warnings = _channel_rows(folder, store)
        warnings.extend(menu_warnings)
        warnings.extend(channel_warnings)
        menu_rows.extend(menu)
        channel_rows.extend(channels)
        covered.append(
            {
                "folder": folder.name,
                "store": store,
                "menu": bool(menu),
                "channels": any(row["kind"] == "total" for row in channels),
                "daily": any(row["kind"] == "daily" for row in channels),
            }
        )
    return {
        "menu_rows": menu_rows,
        "channel_rows": channel_rows,
        "stores": covered,
        "warnings": warnings,
    }


def write_outputs(loaded, menu_dir, channel_dir):
    """Write one menu-mix file and one channel-sales file for the loaded period."""
    menu_dir = Path(menu_dir)
    channel_dir = Path(channel_dir)
    menu_dir.mkdir(parents=True, exist_ok=True)
    channel_dir.mkdir(parents=True, exist_ok=True)
    menu_path = _dated_path(menu_dir, MENU_OUT, loaded["menu_rows"], "period_start", "period_end")
    channel_path = _dated_path(channel_dir, CHANNEL_OUT, loaded["channel_rows"], "period_from", "period_to")
    if menu_path:
        _write_menu(menu_path, loaded["menu_rows"])
    if channel_path:
        _write_channels(channel_path, loaded["channel_rows"])
    return menu_path, channel_path


def _store_name(folder, labels):
    text = ""
    store_file = folder / "store.txt"
    if store_file.exists():
        text = store_file.read_text(encoding="utf-8", errors="replace").strip()
        if text and labels:
            matched = match_menu_store(text, labels)
            if matched:
                return matched, ""
            if text in labels:
                return text, ""
        if text and not labels:
            return text, ""
    region, code = _folder_code(folder.name)
    if code is None:
        return None, "the folder name is not a store code and store.txt did not match."
    hits = []
    for label in labels:
        name_words, label_code = _name_and_code(label)
        if label_code != code or not name_words:
            continue
        label_region = _label_region(label)
        if region and label_region and label_region != region:
            continue
        if region and not label_region:
            continue
        hits.append(label)
    if len(hits) == 1:
        return hits[0], ""
    if text:
        return None, "store.txt did not match one Posist store."
    if len(hits) > 1:
        return None, "the store code matches more than one Posist store."
    return None, "the store code is not on the Posist list."


def _folder_code(name):
    parts = name.split("-")
    if len(parts) == 2 and parts[0] in {"01", "02"} and parts[1].isdigit():
        return parts[0], str(int(parts[1]))
    if name.isdigit():
        return None, str(int(name))
    return None, None


def _label_region(label):
    match = re.search(r"\((01|02)/\d+\)", str(label))
    if match:
        return match.group(1)
    return None


def _menu_rows(folder, store):
    books = []
    warnings = []
    for path in sorted(folder.glob("menu_items_*.xlsx")):
        span = _menu_span(path.name)
        if span is None:
            warnings.append(f"{folder.name}/{path.name} has no period in the name, so it was skipped.")
            continue
        items = _read_menu_book(path)
        if items is None:
            warnings.append(f"{folder.name}/{path.name} has no item sales columns, so it was skipped.")
            continue
        books.append((span, items))
    if not books:
        return [], warnings
    start = min(span[0] for span, _items in books)
    end = max(span[1] for span, _items in books)
    combined = {}
    for _span, items in books:
        for item in items:
            key = re.sub(r"\s+", " ", item["item"]).casefold()
            slot = combined.setdefault(key, {"item": item["item"], "gross": 0.0, "orders": 0.0, "gross_blank": False, "orders_blank": False})
            if item["gross"] is None:
                slot["gross_blank"] = True
            else:
                slot["gross"] += item["gross"]
            if item["orders"] is None:
                slot["orders_blank"] = True
            else:
                slot["orders"] += item["orders"]
    gross_values = [slot["gross"] for slot in combined.values() if not slot["gross_blank"]]
    gross_total = sum(gross_values) if gross_values and not any(slot["gross_blank"] for slot in combined.values()) else None
    rows = []
    for slot in combined.values():
        gross = None if slot["gross_blank"] else slot["gross"]
        orders = None if slot["orders_blank"] else slot["orders"]
        contribution = None
        if gross is not None and gross_total not in (None, 0):
            contribution = gross / gross_total * 100.0
        rows.append(
            {
                "store": store,
                "item": slot["item"],
                "gross": gross,
                "orders": orders,
                "contribution_pct": contribution,
                "period_start": start,
                "period_end": end,
            }
        )
    rows.sort(key=lambda row: (-(row["gross"] if row["gross"] is not None else float("-inf")), row["item"].casefold()))
    return rows, warnings


def _read_menu_book(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return None
        columns = {_header(cell): index for index, cell in enumerate(header) if cell is not None}
        item_at = columns.get("item name")
        sales_at = columns.get("total sales")
        orders_at = columns.get("total orders")
        if item_at is None or sales_at is None or orders_at is None:
            return None
        parsed = []
        for raw in rows:
            if not raw or item_at >= len(raw):
                continue
            item = str(raw[item_at] or "").strip()
            if not item:
                continue
            parsed.append(
                {
                    "item": item,
                    "gross": _cell_number(raw, sales_at),
                    "orders": _cell_number(raw, orders_at),
                }
            )
        return parsed
    finally:
        workbook.close()


def _channel_rows(folder, store):
    rows = []
    warnings = []
    range_files = sorted(folder.glob("source_range_*.tsv"))
    for path in range_files:
        span = _menu_span(path.name.replace("source_range_", "menu_items_"))
        if span is None:
            warnings.append(f"{folder.name}/{path.name} has no period in the name, so it was skipped.")
            continue
        start, end = span
        for record in _read_tsv(path):
            source = str(record.get("source") or "").strip()
            if not source:
                continue
            rows.append(
                {
                    "store": store,
                    "channel": source,
                    "gross": parse_number(record.get("sales")),
                    "orders": parse_number(record.get("orders")),
                    "when": "",
                    "kind": "total",
                    "period_from": start,
                    "period_to": end,
                }
            )
    for path in sorted(folder.glob("source_daily_*.tsv")):
        period = _range_period(folder)
        for record in _read_tsv(path):
            source = str(record.get("source") or "").strip()
            when = str(record.get("date") or "").strip()
            if not source or not when:
                continue
            rows.append(
                {
                    "store": store,
                    "channel": source,
                    "gross": parse_number(record.get("sales")),
                    "orders": parse_number(record.get("orders")),
                    "when": when,
                    "kind": "daily",
                    "period_from": period[0] if period else None,
                    "period_to": period[1] if period else None,
                }
            )
    return rows, warnings


def _range_period(folder):
    for path in sorted(folder.glob("source_range_*.tsv")):
        span = _menu_span(path.name.replace("source_range_", "menu_items_"))
        if span:
            return span
    return None


def _read_tsv(path):
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def _menu_span(name):
    match = _MENU_FILE.search(name)
    if not match:
        return None
    start = parse_date(match.group(1))
    end_text = match.group(2)
    if start is None:
        return None
    if len(end_text) == 10:
        end = parse_date(end_text)
    else:
        month, day = end_text.split("-")
        year = start.year + (1 if int(month) < start.month else 0)
        try:
            end = date(year, int(month), int(day))
        except ValueError:
            return None
    if end is None:
        return None
    return start, end


def _header(cell):
    return re.sub(r"[^a-z0-9]+", " ", str(cell or "").strip().casefold()).strip()


def _cell_number(raw, index):
    if index >= len(raw):
        return None
    return parse_number(raw[index])


def _dated_path(directory, pattern, rows, start_key, end_key):
    starts = [row[start_key] for row in rows if row.get(start_key)]
    ends = [row[end_key] for row in rows if row.get(end_key)]
    if not starts or not ends:
        return None
    start = min(starts).isoformat()
    end = max(ends).isoformat()
    return directory / pattern.format(start=start, end=end)


def _csv_number(value):
    if value is None:
        return ""
    return f"{value:.10f}".rstrip("0").rstrip(".")


def _write_menu(path, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["store", "item", "total_sales", "total_orders", "contribution_pct", "period_start", "period_end"])
        for row in rows:
            writer.writerow(
                [
                    row["store"],
                    row["item"],
                    _csv_number(row["gross"]),
                    _csv_number(row["orders"]),
                    _csv_number(row["contribution_pct"]),
                    row["period_start"].isoformat() if row["period_start"] else "",
                    row["period_end"].isoformat() if row["period_end"] else "",
                ]
            )


def _write_channels(path, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["date", "period_from", "period_to", "store", "channel", "gross", "orders"])
        for row in rows:
            writer.writerow(
                [
                    row["when"],
                    row["period_from"].isoformat() if row["period_from"] else "",
                    row["period_to"].isoformat() if row["period_to"] else "",
                    row["store"],
                    row["channel"],
                    _csv_number(row["gross"]),
                    _csv_number(row["orders"]),
                ]
            )
